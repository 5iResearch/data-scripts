"""
Investor Suite articles (Chart of Interest, Market Musings, Stock of Interest): Word draft -> blog HTML.

Point it at an article folder (e.g. Documents/Investor Suite/2026/Sep/COI) or a whole month folder. For each
article it:
  * picks the Word draft - the one that best matches the month's Investor Suite PDF when there is one,
    otherwise the most recently saved (pin a file with "Document:" in post.txt)
  * past issues, where the month's published PDF is in the folder: takes the wording and every chart from the
    PDF itself (see investor_suite_pdf.py) and uses the draft only for headings and chart positions, so the
    post matches what members received. "Source: word" in post.txt, or --source word, builds from the draft
  * turns the draft into blog markup: bold-only lines become subheadings, pasted charts are saved as
    investor-suite-<yyyy>-<mm>-<type>-fig<N>.png with their "Figure N" / "Source:" lines as captions,
    Word tables become HTML tables, "(DONE)" tags and section numbers are dropped
  * Stock of Interest only: adds the Investment Thesis / Sweet Spot bullets from the companion doc and an
    Investor Considerations box (figures from post.txt, or read from the PDF for past issues)
  * writes the short description that goes with the post: up to three one-line bullets and one embedded
    chart ("Summary 1-3" and "Summary Chart" in post.txt; left blank, the bullets are suggested). The chart
    is shown through the site's [[embedchart:<slug>]] tag, so the kit also gives the slug, title and link to
    register it with
  * writes _build/post.html and _build/kit.html (copy buttons + preview) next to the draft

Nothing is published or committed: these are member posts, so text and images stay in the Documents folder.
The PNGs are uploaded to the site's file manager by hand, keeping their names: the <img> tags already point
at /files/www/<name>.

Usage:  python scripts/build_investor_suite.py "<article or month folder>" [--no-open]
        (or drag the folder onto "Build Investor Suite.bat")
"""

import argparse
import html
import io
import json
import os
import re
import subprocess
import sys
import webbrowser
from datetime import datetime

import docx
from docx.oxml.ns import qn
from docx.table import Table
from PIL import Image

try:
    import fitz  # PyMuPDF, only needed for the published PDF
    import investor_suite_pdf as pdfsrc
except ImportError:
    fitz = pdfsrc = None

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
KIT_TEMPLATE = os.path.join(REPO_ROOT, "templates", "investor_suite_kit.html")
SITE_IMAGE_DIR = "/files/www"
SITE_URL = "https://www.5iresearch.ca"      # an embedded chart is registered on the site with its full link

FONT = "Arial,Helvetica,sans-serif"
BLUE, ORANGE, INK = "#1F79BE", "#C67A29", "#363636"
MUTED, RULE, PANEL = "#6F6F6F", "#E3E6EA", "#F6F8FA"
BLUE_TINT, ORANGE_TINT = "#EEF5FB", "#FCF3EA"
POST_WIDTH = 800           # px, the widest a full-width Word image is shown
SINGLE_WIDTH = 780         # px, the reading column in the "single" layout
COLUMNS_WIDTH = 1120       # px, the whole post in the "columns" layout
COLUMN_MIN = 340           # px, below two of these the "columns" layout drops to one column
LAYOUTS = ("single", "columns", "sections", "sidebar", "wide")
MAX_IMAGE_PX = 1800

KINDS = {"coi": "Chart of Interest", "musings": "Market Musings", "soi": "Stock of Interest"}
MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October",
          "November", "December"]
SKIP_DOC = re.compile(r"notes|wording|memes|methodology", re.I)
SKIP_DIR = re.compile(r"^(_build|old|memes|market memes|claude outputs)$", re.I)
TITLE_PREFIX = re.compile(r"^(COI|Chart of Interest|Market Musings?(?: Draft)?)\s*[–—:_-]*\s*", re.I)
DONE_TAG = re.compile(r"\s*\(\s*DON+E+\s*\)\s*$", re.I)
BULLET = re.compile(r"^[●·•▪◦]\s*")
CAPTION = re.compile(r"^(Figure|Chart|Table|Exhibit)\s+\d+\b.{0,200}$|^Sources?\s*:", re.I)

HOLDINGS_DISCLOSURE = ("*Authors, directors, partners and/or officers of 5i Research do not hold a financial or "
                       "other interest in the above companies at the time of publishing. The i2i Fund does not "
                       "hold a financial or other interest in the above companies at the time of publishing.")
ANALYST_DISCLOSURE = ("Analysts of 5i Research responsible for this report do not have a financial or other "
                      "interest in {t} and will not trade in shares within 48 hours after publishing this "
                      "report. The i2i Fund does not have a financial or other interest in {t}.")
CONSIDERATION_FIELDS = ["Current Price", "Price Target", "Implied Return", "Investor Type/Style",
                        "Dividend Yield", "Confidence"]
SUMMARY_BULLETS = 3
SUMMARY_MAX = 110          # characters; a suggested bullet longer than this is cut at a clause
SUMMARY_HEIGHT = "auto"    # the embed tag's height: the small chart's own height, so nothing is cropped or padded
SUMMARY_WIDTH = "600"      # px, the short description's own smaller copy of the chart
SUMMARY_STYLES = ("bullets", "pill-blue", "pill-orange", "plain")      # the first is the default


# --------------------------------------------------------------------------- folders

def kind_of(name):
    n = name.lower()
    if "thesis" in n or n.startswith("soi") or re.match(r"^[a-z.]{1,6} draft", n):
        return "soi"
    if n.startswith("coi") or "chart of interest" in n:
        return "coi"
    if "musing" in n:
        return "musings"
    return None


def folder_date(folder):
    """(year, month) from the path, e.g. .../2026/Sep/COI or .../2021/Aug 2021/Musings."""
    year = month = None
    for part in reversed(os.path.abspath(folder).split(os.sep)):
        if month is None:
            m = re.match(r"^([A-Za-z]{3})", part)
            if m and m.group(1).title() in [x[:3] for x in MONTHS] and not kind_of(part):
                month = [x[:3] for x in MONTHS].index(m.group(1).title()) + 1
        y = re.search(r"\b(20\d\d)\b", part)
        if y and year is None:
            year = int(y.group(1))
    return year, month


def month_pdf(folder):
    """The month's published Investor Suite PDF: in this folder or the one above it."""
    for d in (folder, os.path.dirname(folder)):
        for f in sorted(os.listdir(d)):
            if f.lower().endswith(".pdf") and "suite" in f.lower():
                return os.path.join(d, f)
    return None


def find_articles(folder):
    """[(folder, kind, [docx names])] for an article folder, or every article under a month folder."""
    found = []
    docs = [f for f in sorted(os.listdir(folder))
            if f.lower().endswith(".docx") and not f.startswith("~$") and not SKIP_DOC.search(f)]
    own = kind_of(os.path.basename(folder))
    by_kind = {}
    for f in docs:
        k = own or kind_of(f)
        if k:
            by_kind.setdefault(k, []).append(f)
    found += [(folder, k, v) for k, v in by_kind.items()]
    if not own:
        for sub in sorted(os.listdir(folder)):
            path = os.path.join(folder, sub)
            if os.path.isdir(path) and not SKIP_DIR.match(sub) and kind_of(sub):
                found += find_articles(path)
    return found


def read_settings(path):
    out = {}
    if os.path.exists(path):
        with open(path, encoding="utf-8-sig") as f:
            for line in f:
                if ":" in line and not line.lstrip().startswith("#"):
                    k, v = line.split(":", 1)
                    out[k.strip().lower()] = v.strip()
    return out


def write_settings(path, kind, values):
    lines = ["# Settings for this post. Edit a value, save, and rebuild. Delete the file to start over.",
             f"Title: {values.get('title', '')}",
             f"Document: {values.get('document', '')}",
             f"Layout: {values.get('layout', LAYOUTS[0])}",
             f"Source: {values.get('source', 'auto')}"]
    if kind == "soi":
        lines += [f"Thesis Document: {values.get('thesis document', '')}",
                  f"Company: {values.get('company', '')}", f"Ticker: {values.get('ticker', '')}"]
        lines += [f"{k}: {values.get(k.lower(), '')}" for k in CONSIDERATION_FIELDS]
    lines.append(f"Disclosure: {values.get('disclosure', '')}")
    lines += summary_settings(values)
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines) + "\n")


def summary_settings(values):
    return (["# Short description shown with the post: up to three short bullets and one chart. Leave the bullets",
             "# blank to have them suggested. Summary Chart is a chart number from the kit. Leave its slug and",
             "# title blank for the standard ones (e.g. coi-2026-09-fig1)."]
            + [f"Summary {i}: {values.get(f'summary {i}', '')}" for i in range(1, SUMMARY_BULLETS + 1)]
            + [f"Summary Chart: {values.get('summary chart', '1')}",
               f"Summary Chart Height: {values.get('summary chart height', SUMMARY_HEIGHT)}",
               f"Summary Chart Slug: {values.get('summary chart slug', '')}",
               f"Summary Chart Title: {values.get('summary chart title', '')}",
               "# The short description uses its own small copy of the chart, this many pixels wide. If the site",
               "# still stretches it across the page, set Summary Chart Canvas to the page's width (e.g. 1200):",
               "# the small chart is then placed at the left of a white image that wide.",
               f"Summary Chart Width: {values.get('summary chart width', SUMMARY_WIDTH)}",
               f"Summary Chart Canvas: {values.get('summary chart canvas', '')}",
               f"# Summary Style: {', '.join(SUMMARY_STYLES)}",
               f"Summary Style: {values.get('summary style', SUMMARY_STYLES[0])}"])


def summary_markup(style, label, bullets):
    """The short description's bullets: plain lines, a bullet list, or a list under a small coloured pill
    naming the kind of article. Kept quiet on purpose: it is a teaser above the article, not a banner."""
    if style == "plain" or not bullets:
        return "\n".join(bullets)
    items = "".join(f'<li style="margin:0 0 4px;">{html.escape(b)}</li>' for b in bullets)
    out = f'<ul style="margin:0 0 12px;padding-left:20px;line-height:1.5;">{items}</ul>'
    if style.startswith("pill"):
        colour = ORANGE if style.endswith("orange") else BLUE
        out = (f'<div style="margin:0 0 10px;"><span style="display:inline-block;background:{colour};color:#ffffff;'
               f'font-family:{FONT};font-size:11px;font-weight:bold;letter-spacing:1px;text-transform:uppercase;'
               f'line-height:1;padding:6px 12px;border-radius:999px;">{html.escape(label)}</span></div>' + out)
    return out


# --------------------------------------------------------------------------- published PDF

def norm(s):
    return re.sub(r"[^a-z0-9]", "", s.lower())


def pdf_pages(path):
    """Each page's text with InDesign's line-end hyphenation undone."""
    if not path or fitz is None:
        return []
    with fitz.open(path) as d:
        return [p.get_text().replace("­\n", "").replace("­", "") for p in d]


def sentences(blocks):
    out = []
    for b in blocks:
        if b["t"] in ("p", "li"):
            out += [s for s in re.split(r"(?<=[.!?])\s+", b["text"]) if len(norm(s)) >= 40]
    return out


def pdf_match(blocks, pages):
    """Share of the draft's sentences that appear in the PDF, the ones that don't, and the PDF pages
    the article sits on."""
    sents = sentences(blocks)
    if not sents or not pages:
        return None, [], []
    page_norm = [norm(p) for p in pages]
    everything = "".join(page_norm)
    hits, missing = [0] * len(pages), []
    for s in sents:
        n = norm(s)
        if n in everything:
            for i, pn in enumerate(page_norm):
                if n in pn:
                    hits[i] += 1
        else:
            missing.append(s)
    on_pages = [i for i, h in enumerate(hits) if h >= 2]
    return 1 - len(missing) / len(sents), missing, on_pages


def pdf_disclosure(pages, on_pages, kind):
    pattern = (r"Analysts of 5i\s*Research responsible.*?The i2i Fund.*?\." if kind == "soi" else
               r"\*?\s*Authors, directors, partners and/or officers of 5i Research.*?time of publishing\."
               r"\s*The i2i Fund.*?time of publishing\.")
    for i in on_pages:
        m = re.search(pattern, re.sub(r"\s+", " ", pages[i]), re.S)
        if m:
            return re.sub(r"\s+", " ", m.group(0)).replace("5iResearch", "5i Research").strip()
    return ""


def pdf_considerations(pages, on_pages):
    for i in on_pages:
        if "Investor Considerations" not in pages[i]:
            continue
        s = re.sub(r"\s+", " ", pages[i].split("Investor Considerations", 1)[1][:600])
        grab = lambda pat: (re.search(pat, s) or [None, ""])[1].strip()
        return {"current price": grab(r"Current Price\s*\$?\s*([\d,]+\.?\d*)"),
                "price target": grab(r"Price Target\s*\$?\s*([\d,]+\.?\d*)"),
                "implied return": grab(r"Implied Return\s*(-?[\d.]+\s*%)"),
                "investor type/style": grab(r"Investor Type/Style\s*(.+?)\s*Dividend Yield"),
                "dividend yield": grab(r"Dividend Yield\s*(.+?)\s*Confidence"),
                "confidence": grab(r"Confidence in Price Target Being\s*Reached\s*([A-Za-z-]+)")}
    return {}


def pdf_company(pages, ticker):
    for p in pages:
        m = re.search(rf"^([A-Z0-9&.,'’ \-]{{3,60}}?)\s*\({re.escape(ticker)}\)\s*$", p, re.M)
        if m:
            return m.group(1).strip().title()
    return ""


# --------------------------------------------------------------------------- Word draft

def _run_html(r):
    text = ""
    for node in r:
        if node.tag == qn("w:t"):
            text += node.text or ""
        elif node.tag == qn("w:tab"):
            text += " "
        elif node.tag in (qn("w:br"), qn("w:cr")):
            text += "\n"
    rpr = r.find(qn("w:rPr"))

    def on(tag):
        el = rpr.find(qn(tag)) if rpr is not None else None
        return el is not None and el.get(qn("w:val")) not in ("0", "false")

    return text, on("w:b"), on("w:i")


def paragraph_lines(para):
    """A paragraph as one entry per line (Shift+Enter breaks split it): html, text, all-bold, all-italic."""
    style_bold = bool(para.style.font.bold)
    lines = [dict(html="", text="", bold=True, italic=True)]

    def walk(el, href=""):
        for child in el:
            if child.tag == qn("w:r"):
                text, b, i = _run_html(child)
                for n, piece in enumerate(text.split("\n")):
                    if n:
                        lines.append(dict(html="", text="", bold=True, italic=True))
                    if not piece:
                        continue
                    s = html.escape(piece, quote=False)
                    if piece.strip():
                        lines[-1]["bold"] &= b or style_bold
                        lines[-1]["italic"] &= i
                    s = f"<em>{s}</em>" if i else s
                    s = f"<strong>{s}</strong>" if b else s
                    if href:
                        s = f'<a href="{html.escape(href)}" style="color:{BLUE};">{s}</a>'
                    lines[-1]["html"] += s
                    lines[-1]["text"] += piece
            elif child.tag == qn("w:hyperlink"):
                rid = child.get(qn("r:id"))
                walk(child, para.part.rels[rid].target_ref if rid and rid in para.part.rels else "")
            elif child.tag in (qn("w:ins"), qn("w:smartTag"), qn("w:sdt"), qn("w:sdtContent")):
                walk(child, href)

    walk(para._element)
    for ln in lines:
        ln["text"] = ln["text"].strip()
        ln["html"] = re.sub(r"</(strong|em)>(\s*)<\1>", r"\2", ln["html"]).strip()
    return [ln for ln in lines if ln["text"]]


def emf_to_png(blob, out_path):
    """Word stores some pasted Excel ranges as EMF, which browsers can't show. Windows can redraw them."""
    src = out_path + ".emf"
    with open(src, "wb") as f:
        f.write(blob)
    quote = lambda v: "'" + v.replace("'", "''") + "'"
    script = (f"Add-Type -AssemblyName System.Drawing; $m = New-Object System.Drawing.Imaging.Metafile({quote(src)}); "
              "$s = 1600.0 / $m.Width; $w = [int]($m.Width * $s); $h = [int]($m.Height * $s); "
              "$b = New-Object System.Drawing.Bitmap($w, $h); $g = [System.Drawing.Graphics]::FromImage($b); "
              "$g.Clear([System.Drawing.Color]::White); $g.InterpolationMode = 'HighQualityBicubic'; "
              f"$g.DrawImage($m, 0, 0, $w, $h); $b.Save({quote(out_path)}, [System.Drawing.Imaging.ImageFormat]::Png); "
              "$g.Dispose(); $b.Dispose(); $m.Dispose()")
    try:
        subprocess.run(["powershell", "-NoProfile", "-Command", script],
                       check=True, capture_output=True, timeout=60)
    finally:
        if os.path.exists(src):
            os.remove(src)


def save_image(blob, ext, crop, out_path):
    """Write the picture as a PNG the way Word shows it (cropped), capped at MAX_IMAGE_PX wide."""
    if ext in (".emf", ".wmf"):
        emf_to_png(blob, out_path)
        with open(out_path, "rb") as f:
            blob = f.read()
    im = Image.open(io.BytesIO(blob))
    im.load()
    if any(crop):
        w, h = im.size
        l, t, r, b = (c / 100000 for c in crop)
        im = im.crop((round(w * l), round(h * t), round(w * (1 - r)), round(h * (1 - b))))
    if ext in (".emf", ".wmf"):        # metafiles often carry a blank, see-through margin; flatten and trim it
        im = im.convert("RGBA")
        im = Image.alpha_composite(Image.new("RGBA", im.size, "white"), im).convert("RGB")
        box = im.convert("L").point(lambda v: 255 if v < 240 else 0).getbbox()
        if box:
            pad = 12
            im = im.crop((max(0, box[0] - pad), max(0, box[1] - pad),
                          min(im.width, box[2] + pad), min(im.height, box[3] + pad)))
    if im.width > MAX_IMAGE_PX:
        im = im.resize((MAX_IMAGE_PX, round(im.height * MAX_IMAGE_PX / im.width)), Image.LANCZOS)
    if im.mode not in ("RGB", "RGBA", "L", "P"):
        im = im.convert("RGB")
    im.save(out_path, "PNG", optimize=True)
    return im.size


def paragraph_images(para, text_width):
    """(blob, ext, crop, share of the page width) for each picture pasted in the paragraph."""
    out = []
    for drawing in para._element.iter(qn("w:drawing")):
        blip = next(drawing.iter(qn("a:blip")), None)
        part = para.part.related_parts.get(blip.get(qn("r:embed"))) if blip is not None else None
        if part is None:
            continue
        extent = next(drawing.iter(qn("wp:extent")), None)
        share = min(1.0, int(extent.get("cx")) / text_width) if extent is not None and text_width else 1.0
        rect = next(drawing.iter(qn("a:srcRect")), None)
        crop = [int(rect.get(k, 0)) if rect is not None else 0 for k in ("l", "t", "r", "b")]
        out.append((part.blob, os.path.splitext(part.partname)[1].lower(), crop, share))
    return out


def table_block(tbl):
    rows = []
    for row in tbl.rows:
        cells, seen = [], set()
        for c in row.cells:            # merged cells repeat; keep each once
            if id(c._tc) not in seen:
                seen.add(id(c._tc))
                cells.append(" ".join(c.text.split()))
        rows.append(cells)
    return dict(t="table", rows=rows, caption=[])


def read_draft(path, kind, img_dir, img_stem, warnings):
    """The draft as a title plus a list of blocks: h, p, li, img, table."""
    d = docx.Document(path)
    sec = d.sections[0]
    text_width = (sec.page_width or 0) - (sec.left_margin or 0) - (sec.right_margin or 0)
    has_h1 = any(p.style.name == "Heading 1" for p in d.paragraphs)
    title, blocks, missing_image = "", [], False

    for el in d.element.body.iterchildren():
        if el.tag == qn("w:tbl"):
            blocks.append(table_block(Table(el, d)))
            continue
        if el.tag != qn("w:p"):
            continue
        para = docx.text.paragraph.Paragraph(el, d)
        for blob, ext, crop, share in paragraph_images(para, text_width):
            n = sum(b["t"] == "img" for b in blocks) + 1
            name = f"{img_stem}-fig{n}.png"
            try:
                w, h = save_image(blob, ext, crop, os.path.join(img_dir, name))
            except Exception as e:
                warnings.append(f"Figure {n} ({ext}) could not be converted: {e}. Save it as a PNG in Word and rebuild.")
                continue
            blocks.append(dict(t="img", name=name, w=w, h=h, share=share, caption=[]))

        style = para.style.name
        centred = para.alignment is not None and int(para.alignment) == 1
        numbered = el.find(qn("w:pPr") + "/" + qn("w:numPr")) is not None
        for ln in paragraph_lines(para):
            text = ln["text"]
            if not title and not blocks and (centred or TITLE_PREFIX.match(text)) and len(text) <= 200:
                title = TITLE_PREFIX.sub("", text).strip() or "?"
                continue
            if text.lower() == "image":
                warnings.append(f'A chart is missing after "{blocks[-1].get("text", "")[:60]}..." - the draft '
                                f'only has the word "image" there.' if blocks else "A chart is missing at the top.")
                blocks.append(dict(t="gap", caption=[]))      # holds the chart's place for the PDF's copy
                missing_image = True
                continue
            is_caption = bool(CAPTION.match(text)) or (ln["italic"] and len(text) <= 400)
            if missing_image and is_caption:
                blocks[-1]["caption"].append(text)
                continue
            missing_image = False
            if blocks and blocks[-1]["t"] in ("img", "table") and is_caption:
                blocks[-1]["caption"].append(text)
                continue
            done = bool(DONE_TAG.search(text))
            short = len(text) <= 140 and not text.endswith((".", ":", ";", ","))
            if style.startswith("Heading") or done or (ln["bold"] and short) or (kind == "soi" and numbered and short
                                                                                and len(text) <= 80):
                clean = re.sub(r"^\d+\.\s+", "", DONE_TAG.sub("", text)).strip()
                blocks.append(dict(t="h", text=clean, level=2 if has_h1 and style != "Heading 1"
                                   and style.startswith("Heading") else 1))
            elif BULLET.match(text) or numbered or "list" in style.lower():
                blocks.append(dict(t="li", text=BULLET.sub("", text),
                                   html=re.sub(r"^((?:<\w+>)*)[●·•▪◦]\s*", r"\1", ln["html"])))
            else:
                blocks.append(dict(t="p", text=text, html=ln["html"]))
    return ("" if title == "?" else title), blocks


def read_thesis(path):
    """{'Investment Thesis': [...], 'The Sweet Spot': [...]} from the companion bullet doc."""
    out, current = {}, None
    for para in docx.Document(path).paragraphs:
        for ln in paragraph_lines(para):
            text = BULLET.sub("", ln["text"]).strip()
            if ln["bold"] and len(text) <= 40:
                current = out.setdefault(text.title(), [])
            elif current is not None and text:
                current.append(text)
    return out


# --------------------------------------------------------------------------- blog markup

def para(inner):
    return f'<p style="font-family:{FONT};font-size:16px;line-height:1.6;color:{INK};margin:0 0 14px;">{inner}</p>'


def bullet_list(items):
    lis = "".join(f'<li style="margin:0 0 6px;">{x}</li>' for x in items)
    return (f'<ul style="margin:0 0 14px;padding-left:22px;font-family:{FONT};font-size:16px;line-height:1.6;'
            f'color:{INK};">{lis}</ul>')


def caption_html(lines):
    if not lines:
        return ""
    text = " &middot; ".join(html.escape(x) for x in lines)
    return (f'<div style="font-family:{FONT};font-size:13px;line-height:1.5;color:{MUTED};text-align:center;'
            f'margin:0 0 20px;"><em>{text}</em></div>')


def figure_html(b, title):
    label = next((c for c in b["caption"] if re.match(r"(Figure|Chart|Table|Exhibit)\s+\d+", c, re.I)), "")
    alt = html.escape(f"{title} - {label}" if label else title)
    width = min(b["w"], round(POST_WIDTH * b["share"])) if b["share"] < 0.85 else None
    size = f"width:100%;max-width:{width}px;" if width else "width:100%;"
    return (f'<div style="margin:20px 0 8px;text-align:center;"><img src="{SITE_IMAGE_DIR}/{b["name"]}" alt="{alt}" '
            f'style="{size}height:auto;display:inline-block;" /></div>' + caption_html(b["caption"]))


def table_html(b):
    cell = f"padding:7px 10px;border:1px solid {RULE};font-family:{FONT};font-size:14px;line-height:1.4;"
    out = []
    for r, row in enumerate(b["rows"]):
        tds = "".join(
            (f'<th style="{cell}background:{BLUE};color:#ffffff;text-align:{"left" if c == 0 else "center"};">'
             if r == 0 else
             f'<td style="{cell}color:{INK};text-align:{"left" if c == 0 else "center"};'
             f'{"background:" + PANEL + ";" if r % 2 == 0 else ""}">')
            + html.escape(v) + ("</th>" if r == 0 else "</td>") for c, v in enumerate(row))
        out.append(f"<tr>{tds}</tr>")
    return ('<div style="overflow-x:auto;margin:18px 0 8px;"><table style="border-collapse:collapse;width:100%;">'
            + "".join(out) + "</table></div>" + caption_html(b["caption"]))


def card(heading, items, tint, accent):
    return (f'<div style="background:{tint};border-left:4px solid {accent};border-radius:4px;padding:16px 18px 6px;'
            f'margin:0 0 16px;"><div style="font-family:{FONT};font-size:13px;font-weight:bold;letter-spacing:1px;'
            f'text-transform:uppercase;color:{accent};margin:0 0 8px;">{html.escape(heading)}</div>'
            + bullet_list([html.escape(x) for x in items]) + "</div>")


def money(v):
    v = v.replace("$", "").strip()
    return f"${v}" if v else ""


def considerations_html(s):
    """The Investor Considerations box. Implied return is worked out from the two prices unless given."""
    implied = s.get("implied return", "")
    try:
        cur, tgt = (float(s[k].replace("$", "").replace(",", "")) for k in ("current price", "price target"))
        implied = implied or f"{(tgt / cur - 1) * 100:.1f}%"
    except (KeyError, ValueError):
        pass
    rows = [("Current Price", money(s.get("current price", ""))), ("Price Target", money(s.get("price target", ""))),
            ("Implied Return", implied), ("Investor Type/Style", s.get("investor type/style", "")),
            ("Dividend Yield", s.get("dividend yield", "")),
            ("Confidence in Price Target Being Reached", s.get("confidence", ""))]
    if not any(v for _, v in rows):
        return ""
    body = "".join(
        f'<div style="padding:9px 14px;border-top:1px solid {RULE};font-family:{FONT};font-size:15px;color:{INK};'
        f'overflow:hidden;"><span style="float:right;font-weight:bold;margin-left:12px;'
        f'{"color:" + BLUE + ";" if k == "Implied Return" else ""}">{html.escape(v) or "&ndash;"}</span>'
        f'{html.escape(k)}</div>' for k, v in rows)
    return (f'<div style="max-width:460px;margin:22px auto;border:1px solid {RULE};border-radius:6px;overflow:hidden;">'
            f'<div style="background:{INK};color:#ffffff;font-family:{FONT};font-size:14px;font-weight:bold;'
            f'letter-spacing:1px;text-transform:uppercase;padding:10px 14px;">Investor Considerations</div>'
            f'{body}</div>')


def build_post(kind, title, blocks, eyebrow, disclosure, thesis=None, considerations="", layout="single"):
    """layout:
      wide      fills the blog's width
      single    one centred reading column
      columns   the text under each heading runs in two columns (one on phones), charts and tables across both
      sections  as columns, with a rule above each section and a line between the columns so each section
                reads as its own block
      sidebar   one text column on the left with that section's charts beside it on the right
    """
    multi = layout in ("columns", "sections")
    head = [f'<div style="font-family:{FONT};font-size:12px;font-weight:bold;letter-spacing:1.5px;'
            f'text-transform:uppercase;color:{ORANGE};margin:0 0 14px;">{html.escape(eyebrow)}</div>']
    cards = [card(heading, items, *((BLUE_TINT, BLUE) if i == 0 else (ORANGE_TINT, ORANGE)))
             for i, (heading, items) in enumerate((thesis or {}).items())]
    if layout in ("columns", "sections", "sidebar") and len(cards) > 1:
        head.append('<div style="display:flex;flex-wrap:wrap;gap:0 16px;">'
                    + "".join(f'<div style="flex:1 1 320px;">{c}</div>' for c in cards) + "</div>")
    else:
        head += cards

    # Group the draft into sections: a heading and what follows it, each item "text", "fig" or "full".
    sections = [dict(h="", items=[])]
    box_at = next((i for i, b in enumerate(blocks) if b["t"] == "h" and "risk" in b["text"].lower()), len(blocks))
    bullets = []

    def flush():
        if bullets:
            sections[-1]["items"].append(("text", bullet_list(bullets[:])))
            bullets.clear()

    for i, b in enumerate(blocks):
        if b["t"] != "li":
            flush()
        if i == box_at and considerations:
            sections[-1]["items"].append(("full", considerations))
        if b["t"] == "h":
            size, colour = (20, BLUE) if b["level"] == 1 else (17, INK)
            top = 10 if layout == "sections" and b["level"] == 1 else 26
            sections.append(dict(level=b["level"], items=[], h=(
                f'<h3 style="margin:{top}px 0 8px;font-family:{FONT};font-size:{size}px;line-height:1.35;'
                f'color:{colour};">{html.escape(b["text"])}</h3>')))
        elif b["t"] == "p":
            sections[-1]["items"].append(("text", para(b["html"])))
        elif b["t"] == "li":
            bullets.append(b["html"])
        elif b["t"] == "img":
            beside = layout == "sidebar" and b["w"] / max(b["h"], 1) < 2.4
            sections[-1]["items"].append(("fig" if beside else "full", figure_html(b, title)))
        elif b["t"] == "table":
            sections[-1]["items"].append(("full", table_html(b)))
    flush()
    if box_at == len(blocks) and considerations:
        sections[-1]["items"].append(("full", considerations))

    def render(items):
        """One section's items: runs of text (and, for sidebar, their charts) between full-width items."""
        out, run = [], []

        def end_run():
            text = "".join(h for k, h in run if k == "text")
            figs = "".join(h for k, h in run if k == "fig")
            if multi and text:
                rule = f"column-rule:1px solid {RULE};" if layout == "sections" else ""
                out.append(f'<div style="column-count:2;column-width:{COLUMN_MIN}px;column-gap:40px;{rule}">'
                           f'{text}</div>')
            elif layout == "sidebar" and figs:
                out.append(f'<div style="display:flex;flex-wrap:wrap;gap:0 36px;align-items:flex-start;">'
                           f'<div style="flex:1 1 420px;">{text}</div><div style="flex:1 1 380px;">{figs}</div></div>')
            elif layout == "sidebar":
                out.append(f'<div style="max-width:{SINGLE_WIDTH}px;">{text}</div>')
            else:
                out.append(text)
            run.clear()

        for k, h in items:
            if k == "full":
                end_run()
                out.append(h)
            else:
                run.append((k, h))
        end_run()
        return "".join(out)

    parts = head
    for s in sections:
        body = s["h"] + render(s["items"])
        if not body:
            continue
        if layout == "sections" and s["h"] and s.get("level") == 1:
            body = f'<div style="border-top:2px solid {RULE};margin:30px 0 0;padding:6px 0 0;">{body}</div>'
        parts.append(body)
    if disclosure:
        parts.append(f'<div style="margin:28px 0 0;padding:12px 0 0;border-top:1px solid {RULE};font-family:{FONT};'
                     f'font-size:12px;line-height:1.5;color:{MUTED};"><em>{html.escape(disclosure)}</em></div>')
    body = "\n".join(parts)
    width = SINGLE_WIDTH if layout == "single" else None if layout == "wide" else COLUMNS_WIDTH
    return f'<div style="max-width:{width}px;margin:0 auto;">\n{body}\n</div>\n' if width else body + "\n"


def suggest_summary(blocks, thesis):
    """Three teaser lines when post.txt has none: the shortest Investment Thesis bullets for a Stock of
    Interest, otherwise the opening sentence of the first, a middle and the last paragraph."""
    points = next(iter(thesis.values()), []) if thesis else []
    if len(points) >= SUMMARY_BULLETS:
        keep = sorted(sorted(range(len(points)), key=lambda i: len(points[i]))[:SUMMARY_BULLETS])
        picks = [points[i] for i in keep]
    else:
        paras = [b["text"] for b in blocks if b["t"] == "p" and len(b["text"]) > 80]
        spots = sorted({0, len(paras) // 2, len(paras) - 1}) if paras else []
        picks = [re.split(r"(?<=[.!?])\s+", paras[i])[0] for i in spots]
        picks = [t if len(t) <= SUMMARY_MAX else t[:t.rfind(" ", 0, SUMMARY_MAX)].rstrip(" ,;") + "…" for t in picks]
    return [t.strip().rstrip(".") for t in picks]


# --------------------------------------------------------------------------- one article

def build_article(folder, kind, docs, open_kit, source=""):
    typed = bool(kind_of(os.path.basename(folder)))
    build_dir = os.path.join(folder, "_build") if typed else os.path.join(folder, "_build", kind)
    settings_path = os.path.join(folder, "post.txt" if typed else f"post-{kind}.txt")
    os.makedirs(build_dir, exist_ok=True)
    for f in os.listdir(build_dir):
        if f.endswith(".png"):
            os.remove(os.path.join(build_dir, f))

    settings, warnings = read_settings(settings_path), []
    year, month = folder_date(folder)
    if not (year and month):
        when = datetime.fromtimestamp(max(os.path.getmtime(os.path.join(folder, f)) for f in docs))
        year, month = year or when.year, month or when.month
        warnings.append(f"Couldn't read the issue month from the folder path; using {MONTHS[month - 1]} {year}.")
    img_stem = f"investor-suite-{year}-{month:02d}-{kind}"
    pdf_path = month_pdf(folder)
    pages = pdf_pages(pdf_path)

    thesis_docs = [f for f in docs if "thesis" in f.lower()]
    drafts = [f for f in docs if f not in thesis_docs]
    if not drafts:
        print(f"  {KINDS[kind]}: no draft found in {folder}")
        return None

    # Pick the draft: pinned in post.txt, else the best match to the PDF, else the latest saved.
    scored = []
    for f in drafts:
        scratch = []
        _, blocks = read_draft(os.path.join(folder, f), kind, build_dir, "_probe", scratch)
        scored.append((pdf_match(blocks, pages)[0], os.path.getmtime(os.path.join(folder, f)), f))
    for f in os.listdir(build_dir):
        if f.startswith("_probe"):
            os.remove(os.path.join(build_dir, f))
    pinned = settings.get("document", "")
    if pinned and pinned not in drafts:
        warnings.append(f'post.txt names "{pinned}", which is not in the folder; picked a draft automatically.')
        pinned = ""
    chosen = pinned or max(scored, key=lambda s: (s[0] if s[0] is not None else -1, s[1]))[2]
    others = [f"{f} ({'%.0f%% match' % (s * 100) if s is not None else 'not compared'})"
              for s, _, f in sorted(scored, key=lambda s: -s[1]) if f != chosen]

    draft_warnings = []
    title, blocks = read_draft(os.path.join(folder, chosen), kind, build_dir, "_draft", draft_warnings)
    score, missing, on_pages = pdf_match(blocks, pages)

    # Past issue with its PDF: the PDF supplies the words and charts, the draft only the structure.
    source = (source or settings.get("source") or "auto").lower()
    article, notes, unused = None, [], []
    if source != "word" and pages and pdfsrc:
        try:
            article = pdfsrc.read_article(pdf_path, kind, blocks, title, build_dir, build_dir, img_stem, on_pages)
        except Exception as e:
            warnings.append(f"Couldn't read this article from the PDF ({e}); built from the Word draft instead.")
        if article is None and source == "pdf":
            warnings.append("This article wasn't found in the PDF, so it was built from the Word draft.")
    for f in os.listdir(build_dir):
        if f.startswith("_draft"):
            if article:
                os.remove(os.path.join(build_dir, f))
            else:
                os.replace(os.path.join(build_dir, f), os.path.join(build_dir, img_stem + f[len("_draft"):]))
    if article:
        first, last = article["pages"][0] + 1, article["pages"][-1] + 1
        notes = [f"Text and charts are taken from the PDF, page{'s' if last > first else ''} {first}"
                 + (f" to {last}" if last > first else "") + "."] + article["notes"]
        if score is not None:
            notes.append(f"The Word draft has {score:.0%} of its sentences word for word in the PDF; where they "
                         "differ, the PDF wording is used.")
        title, blocks, on_pages, unused = article["title"] or title, article["blocks"], article["pages"], article["unused"]
    else:
        warnings += draft_warnings
        for b in blocks:
            if b["t"] == "img":
                b["name"] = img_stem + b["name"][len("_draft"):]
    if len(drafts) > 1 and not pinned and score is None:
        warnings.append("Several drafts and no PDF to compare against, so the most recently saved was used. "
                        'Set "Document:" in post.txt to pin the right one.')

    thesis, considerations = {}, ""
    if kind == "soi":
        ticker = settings.get("ticker") or (re.match(r"^([A-Z.]{1,6}) ", chosen) or [None, ""])[1]
        company = settings.get("company") or pdf_company(pages, ticker) if ticker else settings.get("company", "")
        if not company and ticker:
            m = re.search(rf"((?:[A-Z][\w&.'’-]*\s+){{1,4}})\({re.escape(ticker)}\)",
                          " ".join(b["text"] for b in blocks if b["t"] == "p"))
            company = m.group(1).strip() if m else ""
        title = f"{company} ({ticker})" if company and ticker else (company or ticker or title)
        thesis_doc = settings.get("thesis document") or (thesis_docs[-1] if thesis_docs else "")
        if article and article["thesis"]:
            thesis = article["thesis"]
        elif thesis_doc and os.path.exists(os.path.join(folder, thesis_doc)):
            thesis = read_thesis(os.path.join(folder, thesis_doc))
        else:
            warnings.append("No Investment Thesis / Sweet Spot document found, so those boxes are left out.")
        from_pdf = pdf_considerations(pages, on_pages)
        for k in CONSIDERATION_FIELDS:
            settings.setdefault(k.lower(), from_pdf.get(k.lower(), ""))
            if not settings[k.lower()] and k != "Implied Return":
                settings[k.lower()] = from_pdf.get(k.lower(), "")
        considerations = considerations_html(settings)
        empty = [k for k in CONSIDERATION_FIELDS if not settings.get(k.lower()) and k != "Implied Return"]
        if empty:
            warnings.append(f"Investor Considerations: fill in {', '.join(empty)} in {os.path.basename(settings_path)} "
                            "and rebuild.")
        settings.update({"ticker": ticker, "company": company, "thesis document": thesis_doc})
    title = settings.get("title") or title or TITLE_PREFIX.sub("", os.path.splitext(chosen)[0]).strip()

    disclosure = settings.get("disclosure") or pdf_disclosure(pages, on_pages, kind)
    if not disclosure:
        disclosure = ANALYST_DISCLOSURE.format(t=settings.get("ticker") or "the company") if kind == "soi" \
            else HOLDINGS_DISCLOSURE
        warnings.append(f"Disclosure: used the standard no-holdings wording. Check it, and edit it in "
                        f"{os.path.basename(settings_path)} if anyone holds a name in this post.")

    if not os.path.exists(settings_path):
        write_settings(settings_path, kind, {**settings, "title": title, "document": chosen, "disclosure": disclosure})
    else:                                       # a post.txt from before these settings were added
        new = [ln for ln in summary_settings(settings)
               if not ln.startswith("#") and ln.split(":", 1)[0].lower() not in settings]
        if new:
            with open(settings_path, "a", encoding="utf-8") as f:
                f.write("\n".join(new) + "\n")
    if score is not None and score < 0.9 and not article:
        warnings.append(f"Only {score:.0%} of this draft's sentences appear in the published PDF, so the text was "
                        "probably edited in InDesign. See the list at the bottom.")
    images = [b for b in blocks if b["t"] == "img"]
    if not images:
        warnings.append("No charts found in the draft.")

    # The short description: a few one-line bullets and one chart, as the site's embed tag.
    bullets = [settings.get(f"summary {i}", "") for i in range(1, SUMMARY_BULLETS + 1)]
    bullets = [b.lstrip("-•· ").strip() for b in bullets if b.strip()]
    suggested = not bullets
    if suggested:
        bullets = suggest_summary(blocks, thesis)
        if not thesis:
            warnings.append("Short description: the three bullets are just sentences lifted from the article. "
                            f"Write your own as Summary 1 to 3 in {os.path.basename(settings_path)} and rebuild.")
    pick = settings.get("summary chart", "").strip() or "1"
    number = int(pick) if pick.isdigit() else 0
    chosen_image = images[number - 1] if 0 < number <= len(images) else None
    if not chosen_image and images:
        warnings.append(f'Short description: "Summary Chart: {pick}" is not one of charts 1 to {len(images)}, so '
                        "no chart was embedded.")
    # The site shows a chart through a registered slug, not a file name: slug + title + the uploaded file's link.
    embed = chart_title = chart_link = ""
    preview = None
    if chosen_image:
        embed = settings.get("summary chart slug", "").strip() or f"{kind}-{year}-{month:02d}-fig{number}"
        chart_title = (settings.get("summary chart title", "").strip()
                       or f"{KINDS[kind]}: {title} ({MONTHS[month - 1][:3]} {year})")
        # Its own small copy: the embed shows a picture as big as the file is, so the file is made small.
        number_of = lambda key, default: int(v) if (v := settings.get(key, "").strip()).isdigit() else default
        width, canvas = number_of("summary chart width", int(SUMMARY_WIDTH)), number_of("summary chart canvas", 0)
        with Image.open(os.path.join(build_dir, chosen_image["name"])) as im:
            small = im.convert("RGB")
        if small.width > width:
            small = small.resize((width, round(small.height * width / small.width)), Image.LANCZOS)
        if canvas > small.width:
            sheet = Image.new("RGB", (canvas, small.height), "white")
            sheet.paste(small, (0, 0))
            small = sheet
        preview = dict(name=os.path.splitext(chosen_image["name"])[0] + "-preview.png", w=small.width,
                       h=small.height, caption=["short description"])
        small.save(os.path.join(build_dir, preview["name"]), "PNG", optimize=True)
        chart_link = f"{SITE_URL}{SITE_IMAGE_DIR}/{preview['name']}"
    height = settings.get("summary chart height", "").strip() or SUMMARY_HEIGHT
    if not height.isdigit():
        height = str(preview["h"]) if preview else "420"
    style = settings.get("summary style", "").strip().lower() or SUMMARY_STYLES[0]
    if style not in SUMMARY_STYLES:
        warnings.append(f'Summary Style "{style}" is not one of {", ".join(SUMMARY_STYLES)}; used {SUMMARY_STYLES[0]}.')
        style = SUMMARY_STYLES[0]
    summary_look = summary_markup(style, KINDS[kind], bullets)
    summary = "\n".join(filter(None, [summary_look, f"[[embedchart:{embed} height={height}]]" if embed else ""]))
    with open(os.path.join(build_dir, "summary.txt"), "w", encoding="utf-8") as f:
        f.write(summary + "\n")

    month_label = f"{MONTHS[month - 1]} {year}"
    post_title = f"{KINDS[kind]}: {title}"
    layout = (os.environ.get("INVESTOR_SUITE_LAYOUT") or settings.get("layout") or LAYOUTS[0]).lower()
    if layout not in LAYOUTS:
        warnings.append(f'Layout "{layout}" is not one of {", ".join(LAYOUTS)}; used {LAYOUTS[0]}.')
        layout = LAYOUTS[0]
    post_html = build_post(kind, title, blocks, f"{KINDS[kind]} · {month_label}", disclosure, thesis,
                           considerations, layout)
    with open(os.path.join(build_dir, "post.html"), "w", encoding="utf-8") as f:
        f.write(post_html)
    data = {"title": post_title, "kind": KINDS[kind], "month": month_label, "post": post_html,
            "warnings": warnings, "document": chosen, "others": others,
            "pdf": os.path.basename(pdf_path) if pages else "", "match": None if score is None else round(score * 100),
            "missing": [] if article else missing[:40], "settings": os.path.basename(settings_path),
            "source": "pdf" if article else "word", "notes": notes, "unused": unused[:60],
            "summary": summary, "summary_bullets": bullets, "summary_suggested": suggested,
            "summary_style": style, "summary_look": summary_look,
            "summary_image": preview["name"] if preview else "", "summary_height": height,
            "chart_slug": embed, "chart_title": chart_title, "chart_link": chart_link,
            "words": sum(len(b["text"].split()) for b in blocks if b["t"] in ("p", "li")),
            "images": [{"name": b["name"], "w": b["w"], "h": b["h"], "caption": " · ".join(b["caption"])}
                       for b in images + ([preview] if preview else [])]}
    with open(KIT_TEMPLATE, encoding="utf-8") as f:
        page = f.read()
    kit_path = os.path.join(build_dir, "kit.html")
    with open(kit_path, "w", encoding="utf-8") as f:
        f.write(page.replace("%%DATA_JSON%%", json.dumps(data).replace("</", "<\\/"))
                    .replace("%%TITLE%%", html.escape(post_title)))

    print(f"\n{post_title}  ({month_label})")
    print(f"  Draft: {chosen}" + (f"  [{score:.0%} match to the PDF]" if score is not None else ""))
    for note in notes:
        print(f"  {note}")
    if unused:
        print(f"  {len(unused)} line(s) of PDF text on those pages were not used (listed in the kit).")
    print(f"  {data['words']} words, {len(images)} chart(s)")
    for w in warnings:
        print(f"  ! {w}")
    print(f"  Kit: {kit_path}")
    if open_kit:
        webbrowser.open("file:///" + kit_path.replace(os.sep, "/"))
    return kit_path


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("folder", help="Article folder (COI, Musings, SOI) or a month folder holding them")
    ap.add_argument("--no-open", action="store_true")
    ap.add_argument("--source", choices=["auto", "pdf", "word"], default="",
                    help="pdf: wording and charts from the month's published PDF (the default when it is there); "
                         "word: build from the draft")
    args = ap.parse_args()

    folder = os.path.abspath(args.folder.strip('"'))
    if not os.path.isdir(folder):
        sys.exit(f"Not a folder: {folder}")
    articles = find_articles(folder)
    if not articles:
        sys.exit("No Chart of Interest, Market Musings or Stock of Interest drafts found. Name the folder COI, "
                 'Musings or SOI, or start the file name with "COI" / "Market Musings".')
    for path, kind, docs in articles:
        build_article(path, kind, docs, not args.no_open, args.source)


if __name__ == "__main__":
    main()
