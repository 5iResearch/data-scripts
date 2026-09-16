"""
Weekly "Benchmark Beaters" blog + email kit.

Reads the most recent Benchmark_Beaters_<date>.xlsx written by
generate_benchmark_beaters.py (no new screening / network calls) and writes a
ready-to-paste kit to outputs/benchmark-beaters-weekly/<date>/:

  teaser.png   - top names from each table in 5i colours, rest of the list
                 teased as "+N more in the full report"
  post.html    - blog body: intro, teaser image, PDF buttons, top-name HTML
                 tables (real text, for SEO) and disclosure
  email.html   - full Mailchimp "code your own" email (inline styles, merge
                 tags for preview text / unsubscribe / address)
  meta.json    - blog title, email subject, preview text, headline stats
  kit.html     - open this each week: paste the two PDF links once, then
                 one-click copy of title / body / subject / email HTML

The two PDFs are downloaded and hosted by hand, so their links are passed in
with --table-url / --charts-url. Without them the HTML keeps the
{{TABLE_PDF_URL}} / {{CHARTS_PDF_URL}} placeholders to replace before sending.
"""

import argparse
import glob
import html
import json
import os
import re
import shutil
from collections import Counter
from datetime import datetime

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
from openpyxl import load_workbook

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BB_DIR = os.path.join(REPO_ROOT, "outputs", "benchmark-beaters")
OUTPUT_ROOT = os.path.join(REPO_ROOT, "outputs", "benchmark-beaters-weekly")
LOGO_PATH = os.path.join(REPO_ROOT, "assets", "Logo_Transparent_1200px.png")
KIT_TEMPLATE = os.path.join(REPO_ROOT, "templates", "bb_weekly_kit.html")
PAGES_BASE = "https://5iresearch.github.io/data-scripts/outputs/benchmark-beaters-weekly"

TRIAL_URL = "https://www.5iresearch.ca/bb"
TABLE_PLACEHOLDER = "{{TABLE_PDF_URL}}"
CHARTS_PLACEHOLDER = "{{CHARTS_PDF_URL}}"

# 5i brand colours
BLUE = "#1F79BE"
ORANGE = "#C67A29"
INK = "#363636"
MUTED = "#7A7A7A"
RULE = "#E3E3E3"
BAND = "#F6F8FA"
BADGE = {"NEW": "#2E86DE", "RPT": "#8E44AD", "STRK": "#C0392B"}

SECTOR_SHORT = {
    "Information Technology": "Info Tech",
    "Communication Services": "Comm Services",
    "Consumer Discretionary": "Cons Discretionary",
    "Consumer Staples": "Cons Staples",
}

DATA_START_ROW = 8
TOP_N = 5
DISCLOSURE = (
    "Disclosure: This list is a quantitative screen and does not constitute any formal opinion "
    "on the names presented. Authors may own shares in non-Canadian securities listed above."
)


# --------------------------------------------------------------------------- data

def latest_workbook(date_str=None):
    if date_str:
        path = os.path.join(BB_DIR, f"Benchmark_Beaters_{date_str}.xlsx")
        if not os.path.exists(path):
            raise FileNotFoundError(path)
        return path, date_str
    dated = sorted(glob.glob(os.path.join(BB_DIR, "Benchmark_Beaters_????-??-??.xlsx")))
    if not dated:
        raise FileNotFoundError(f"No dated Benchmark Beaters workbooks in {BB_DIR}")
    path = dated[-1]
    return path, re.search(r"(\d{4}-\d{2}-\d{2})", path).group(1)


def read_table(ws):
    rows = []
    for signal, ticker, company, sector, industry, ret, mcap, desc in ws.iter_rows(
        min_row=DATA_START_ROW, min_col=2, max_col=9, values_only=True
    ):
        if not ticker or not isinstance(ret, (int, float)):
            break
        rows.append({
            "signal": signal or "", "ticker": str(ticker), "company": company or "",
            "sector": sector or "", "industry": industry or "", "ret": float(ret),
            "mcap": mcap if isinstance(mcap, (int, float)) else None, "desc": desc or "",
        })
    rows.sort(key=lambda r: r["ret"], reverse=True)
    return rows


def summarize(rows):
    sectors = Counter(r["sector"] for r in rows if r["sector"])
    top_sector, top_sector_n = sectors.most_common(1)[0] if sectors else ("", 0)
    return {
        "count": len(rows),
        "new": sum(r["signal"] == "NEW" for r in rows),
        "streak": sum(r["signal"] == "STRK" for r in rows),
        "top_sector": top_sector,
        "top_sector_n": top_sector_n,
        "leader": rows[0] if rows else None,
    }


# --------------------------------------------------------------------------- copy

def pct(x):
    return f"{x * 100:+.0f}%"


def long_date(d):
    return f"{d:%B} {d.day}, {d.year}"


def market_sentence(label, bench, s):
    if not s["count"]:
        return f"No {label} stocks made a fresh six-month high relative to the {bench} this week."
    lead = s["leader"]
    parts = [
        f"In {label}, {s['count']} stocks made a fresh six-month high relative to the {bench}"
        + (f", {s['new']} of them new to the list this week." if s["new"] else ".")
    ]
    parts.append(
        f"{lead['company']} ({lead['ticker']}) leads with a {pct(lead['ret'])} six-month return"
        + (f", and {s['top_sector']} is the most represented sector with {s['top_sector_n']} names."
           if s["top_sector_n"] > 1 else ".")
    )
    return " ".join(parts)


def build_copy(cdn, us, as_of):
    c, u = summarize(cdn), summarize(us)
    total = c["count"] + u["count"]
    leaders = [s["leader"]["ticker"] for s in (c, u) if s["leader"]]
    intro = [
        f"This week's Benchmark Beaters screen found {total} Canadian and U.S. stocks outperforming "
        f"their benchmarks, each hitting a new six-month relative high in the last five trading days.",
        market_sentence("Canada", "TSX Composite", c),
        market_sentence("the U.S.", "S&P 500", u),
        "Download the full table for every name, and the chart pack to see how each one's "
        "relative strength has built over six months and ten years.",
    ]
    return {
        "as_of": as_of.strftime("%Y-%m-%d"),
        "blog_title": f"Benchmark Beaters Weekly: {total} Stocks Beating Their Benchmarks ({long_date(as_of)})",
        "email_subject": f"Benchmark Beaters: {' & '.join(leaders)} lead {total} stocks beating the market",
        "preview_text": (f"{c['count']} Canadian and {u['count']} U.S. names at six-month relative highs. "
                         f"Full table and charts inside."),
        "intro": intro,
        "stats": {"cdn": {k: v for k, v in c.items() if k != "leader"},
                  "us": {k: v for k, v in u.items() if k != "leader"}},
    }


# --------------------------------------------------------------------------- teaser png

def _draw_panel(fig, top, height, title, bench, rows, color):
    ax = fig.add_axes([0.05, top - height, 0.90, height])
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")
    shown = rows[:TOP_N]
    n_slots = TOP_N + 2  # header + rows + "more" line
    slot = 1 / (n_slots + 0.6)

    ax.text(0, 1, title, fontsize=19, fontweight="bold", color=INK, va="top")
    ax.text(1, 1, f"{len(rows)} names  ·  vs. {bench}", fontsize=12, color=MUTED, va="top", ha="right")
    y = 1 - slot * 0.95
    ax.plot([0, 1], [y, y], color=color, lw=2.5)
    ax.text(0.10, y - slot * 0.35, "TICKER", fontsize=9, color=MUTED, fontweight="bold", va="center")
    ax.text(0.22, y - slot * 0.35, "COMPANY", fontsize=9, color=MUTED, fontweight="bold", va="center")
    ax.text(0.60, y - slot * 0.35, "SECTOR", fontsize=9, color=MUTED, fontweight="bold", va="center")
    ax.text(1.0, y - slot * 0.35, "6-MO RETURN", fontsize=9, color=MUTED, fontweight="bold",
            va="center", ha="right")
    y -= slot * 0.75

    max_ret = max((r["ret"] for r in shown), default=1) or 1
    for i, r in enumerate(shown):
        cy = y - slot * (i + 0.5)
        if i % 2 == 0:
            ax.add_patch(plt.Rectangle((0, cy - slot / 2), 1, slot, color=BAND, lw=0))
        badge = BADGE.get(r["signal"])
        if badge:
            ax.add_patch(FancyBboxPatch((0.008, cy - slot * 0.22), 0.065, slot * 0.44,
                                        boxstyle="round,pad=0,rounding_size=0.008", color=badge, lw=0))
            ax.text(0.0405, cy, r["signal"], fontsize=8.5, color="white", fontweight="bold",
                    ha="center", va="center")
        ax.text(0.10, cy, r["ticker"], fontsize=13, fontweight="bold", color=color, va="center")
        company = r["company"] if len(r["company"]) <= 30 else r["company"][:29] + "…"
        ax.text(0.22, cy, company, fontsize=12, color=INK, va="center")
        sector = SECTOR_SHORT.get(r["sector"], r["sector"])
        ax.text(0.60, cy, sector, fontsize=10.5, color=MUTED, va="center")
        bar_w = 0.11 * max(r["ret"], 0) / max_ret
        ax.add_patch(plt.Rectangle((0.78, cy - slot * 0.16), bar_w, slot * 0.32, color=color, lw=0, alpha=0.85))
        ax.text(1.0, cy, pct(r["ret"]), fontsize=12.5, fontweight="bold", color=INK, va="center", ha="right")

    more = len(rows) - len(shown)
    my = y - slot * (len(shown) + 0.55)
    msg = (f"+ {more} more {title.split()[0]} names in the full report" if more > 0
           else "Full details in the report")
    ax.text(0.5, my, msg, fontsize=12, color=ORANGE, fontweight="bold", ha="center", va="center",
            style="italic")


def render_teaser(cdn, us, as_of, path):
    fig = plt.figure(figsize=(12, 13.2), dpi=100)
    fig.patch.set_facecolor("white")

    if os.path.exists(LOGO_PATH):
        logo = plt.imread(LOGO_PATH)
        h, w = logo.shape[:2]
        lw = 0.16
        lh = lw * (h / w) * (12 / 13.2)
        lax = fig.add_axes([0.05, 0.975 - lh, lw, lh])
        lax.imshow(logo)
        lax.axis("off")
    fig.text(0.95, 0.965, "BENCHMARK BEATERS", fontsize=26, fontweight="bold", color=INK, ha="right", va="top")
    fig.text(0.95, 0.928, f"Stocks at 6-month relative highs  ·  Week of {long_date(as_of)}",
             fontsize=13, color=MUTED, ha="right", va="top")
    fig.add_artist(plt.Line2D([0.05, 0.95], [0.905, 0.905], color=ORANGE, lw=3))

    _draw_panel(fig, 0.875, 0.38, "Canadian Stocks", "TSX Composite", cdn, BLUE)
    _draw_panel(fig, 0.465, 0.38, "U.S. Stocks", "S&P 500", us, ORANGE)

    fig.add_artist(plt.Line2D([0.05, 0.95], [0.058, 0.058], color=RULE, lw=1))
    fig.text(0.05, 0.035, "NEW = new this week   RPT = 2–3 weeks   STRK = 4+ weeks", fontsize=9.5,
             color=MUTED, va="center")
    fig.text(0.95, 0.035, "Source: 5i Research, Koyfin, Yahoo Finance. Not a recommendation.",
             fontsize=9.5, color=MUTED, va="center", ha="right")
    fig.savefig(path, facecolor="white", dpi=200)
    plt.close(fig)


# --------------------------------------------------------------------------- html

FONT = "Arial,Helvetica,sans-serif"


def _button(url, label, bg):
    return (f'<a href="{html.escape(url)}" style="display:inline-block;background:{bg};color:#ffffff;'
            f'font-family:{FONT};font-size:15px;font-weight:bold;text-decoration:none;'
            f'padding:12px 22px;border-radius:4px;margin:4px 6px;">{label}</a>')


def _table(rows, title, color):
    head_cells = "".join(
        f'<th style="text-align:{a};padding:8px 10px;border-bottom:2px solid {color};font-size:12px;'
        f'color:{MUTED};text-transform:uppercase;">{h}</th>'
        for h, a in (("Ticker", "left"), ("Company", "left"), ("Sector", "left"), ("6-Mo Return", "right"))
    )
    body = []
    for i, r in enumerate(rows[:TOP_N]):
        bg = BAND if i % 2 == 0 else "#ffffff"
        body.append(
            f'<tr style="background:{bg};">'
            f'<td style="padding:8px 10px;font-weight:bold;color:{color};">{html.escape(r["ticker"])}</td>'
            f'<td style="padding:8px 10px;color:{INK};">{html.escape(r["company"])}</td>'
            f'<td style="padding:8px 10px;color:{MUTED};">{html.escape(r["sector"])}</td>'
            f'<td style="padding:8px 10px;text-align:right;font-weight:bold;color:{INK};">{pct(r["ret"])}</td>'
            f'</tr>'
        )
    more = len(rows) - min(len(rows), TOP_N)
    more_line = (f'<p style="margin:6px 0 0;font-size:13px;color:{ORANGE};font-style:italic;">'
                 f'+ {more} more in the full report</p>' if more > 0 else "")
    return (
        f'<h3 style="font-family:{FONT};color:{INK};margin:24px 0 8px;">{title}</h3>'
        f'<table role="presentation" width="100%" cellpadding="0" cellspacing="0" '
        f'style="border-collapse:collapse;font-family:{FONT};font-size:14px;">'
        f'<thead><tr>{head_cells}</tr></thead><tbody>{"".join(body)}</tbody></table>{more_line}'
    )


def _body(copy, cdn, us, img_src, table_url, charts_url):
    paras = "".join(
        f'<p style="font-family:{FONT};font-size:16px;line-height:1.55;color:{INK};margin:0 0 14px;">'
        f'{html.escape(p)}</p>' for p in copy["intro"]
    )
    buttons = (f'<p style="text-align:center;margin:22px 0;">'
               f'{_button(table_url, "Download the Full Table (PDF)", BLUE)}'
               f'{_button(charts_url, "Download the Chart Pack (PDF)", ORANGE)}</p>')
    return (
        paras
        + f'<p style="margin:18px 0;"><a href="{html.escape(table_url)}"><img src="{html.escape(img_src)}" '
          f'alt="Benchmark Beaters: top Canadian and U.S. stocks this week" width="600" '
          f'style="width:100%;max-width:600px;height:auto;border:0;display:block;margin:0 auto;"></a></p>'
        + buttons
        + _table(cdn, "Top Canadian Benchmark Beaters", BLUE)
        + _table(us, "Top U.S. Benchmark Beaters", ORANGE)
        + buttons
        + f'<p style="font-family:{FONT};font-size:16px;line-height:1.55;color:{INK};margin:22px 0 14px;">'
          f'Want to know which of these we would actually buy? '
          f'<a href="{TRIAL_URL}" style="color:{BLUE};font-weight:bold;">Start your 14-day free trial of 5i Research</a>.</p>'
        + f'<p style="font-family:{FONT};font-size:12px;line-height:1.5;color:{MUTED};font-style:italic;">'
          f'{html.escape(DISCLOSURE)}</p>'
    )


def build_post(copy, cdn, us, img_src, table_url, charts_url):
    return f'<div class="bb-weekly">\n{_body(copy, cdn, us, img_src, table_url, charts_url)}\n</div>\n'


def build_kit_page(copy, post_html, email_html, table_url, charts_url, n_cdn, n_us):
    with open(KIT_TEMPLATE, encoding="utf-8") as f:
        page = f.read()
    data = {"blog_title": copy["blog_title"], "email_subject": copy["email_subject"],
            "preview_text": copy["preview_text"], "post": post_html, "email": email_html,
            "placeholders": [TABLE_PLACEHOLDER, CHARTS_PLACEHOLDER]}
    # "</" inside a <script> block would end it early
    data_json = json.dumps(data).replace("</", "<\\/")
    keep = lambda u: "" if u in (TABLE_PLACEHOLDER, CHARTS_PLACEHOLDER) else html.escape(u)
    for key, val in {"AS_OF": copy["as_of"], "N_CDN": str(n_cdn), "N_US": str(n_us),
                     "TABLE_URL": keep(table_url), "CHARTS_URL": keep(charts_url),
                     "DATA_JSON": data_json}.items():
        page = page.replace(f"%%{key}%%", val)
    return page


def build_email(copy, cdn, us, img_src, table_url, charts_url):
    return f"""<!DOCTYPE html>
<html><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>*|MC:SUBJECT|*</title></head>
<body style="margin:0;padding:0;background:#F2F4F6;">
<span style="display:none;max-height:0;overflow:hidden;">*|MC_PREVIEW_TEXT|*</span>
<table role="presentation" width="100%" cellpadding="0" cellspacing="0" style="background:#F2F4F6;">
<tr><td align="center" style="padding:24px 12px;">
<table role="presentation" width="640" cellpadding="0" cellspacing="0" style="max-width:640px;width:100%;background:#ffffff;">
<tr><td style="padding:18px 20px;border-bottom:3px solid {ORANGE};font-family:{FONT};">
<span style="font-size:22px;font-weight:bold;color:{INK};">Benchmark Beaters Weekly</span><br>
<span style="font-size:13px;color:{MUTED};">{html.escape(long_date(datetime.strptime(copy['as_of'], '%Y-%m-%d')))} &middot; 5i Research</span>
</td></tr>
<tr><td style="padding:22px 20px;">
{_body(copy, cdn, us, img_src, table_url, charts_url)}
</td></tr>
<tr><td style="padding:16px 20px;background:#F6F8FA;font-family:{FONT};font-size:11px;line-height:1.5;color:{MUTED};text-align:center;">
You're receiving this because you subscribed to Benchmark Beaters at 5iresearch.ca.<br>
*|LIST:ADDRESSLINE|*<br>
<a href="*|UPDATE_PROFILE|*" style="color:{MUTED};">Update preferences</a> &middot;
<a href="*|UNSUB|*" style="color:{MUTED};">Unsubscribe</a>
</td></tr>
</table>
</td></tr></table>
</body></html>
"""


# --------------------------------------------------------------------------- main

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--date", help="Workbook date YYYY-MM-DD (default: latest)")
    ap.add_argument("--table-url", default=TABLE_PLACEHOLDER, help="Link for the Benchmark Beaters PDF")
    ap.add_argument("--charts-url", default=CHARTS_PLACEHOLDER, help="Link for the chart pack PDF")
    ap.add_argument("--image-url", help="Hosted teaser.png URL (default: GitHub Pages copy)")
    args = ap.parse_args()

    xlsx, date_str = latest_workbook(args.date)
    as_of = datetime.strptime(date_str, "%Y-%m-%d")
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    cdn, us = read_table(wb["Cdn Hardcoded"]), read_table(wb["US Hardcoded"])
    wb.close()
    print(f"{os.path.basename(xlsx)}: {len(cdn)} Cdn, {len(us)} US")

    out_dir = os.path.join(OUTPUT_ROOT, date_str)
    os.makedirs(out_dir, exist_ok=True)
    copy = build_copy(cdn, us, as_of)
    img_src = args.image_url or f"{PAGES_BASE}/{date_str}/teaser.png"

    render_teaser(cdn, us, as_of, os.path.join(out_dir, "teaser.png"))
    # built with placeholders so kit.html can swap in the links typed there
    post_tpl = build_post(copy, cdn, us, img_src, TABLE_PLACEHOLDER, CHARTS_PLACEHOLDER)
    email_tpl = build_email(copy, cdn, us, img_src, TABLE_PLACEHOLDER, CHARTS_PLACEHOLDER)

    def with_links(s):
        return (s.replace(TABLE_PLACEHOLDER, html.escape(args.table_url))
                 .replace(CHARTS_PLACEHOLDER, html.escape(args.charts_url)))

    with open(os.path.join(out_dir, "post.html"), "w", encoding="utf-8") as f:
        f.write(with_links(post_tpl))
    with open(os.path.join(out_dir, "email.html"), "w", encoding="utf-8") as f:
        f.write(with_links(email_tpl))
    with open(os.path.join(out_dir, "kit.html"), "w", encoding="utf-8") as f:
        f.write(build_kit_page(copy, post_tpl, email_tpl, args.table_url, args.charts_url, len(cdn), len(us)))
    with open(os.path.join(out_dir, "meta.json"), "w", encoding="utf-8") as f:
        json.dump({**copy, "workbook": os.path.basename(xlsx), "image_url": img_src,
                   "table_url": args.table_url, "charts_url": args.charts_url}, f, indent=2)

    latest_dir = os.path.join(OUTPUT_ROOT, "latest")
    shutil.rmtree(latest_dir, ignore_errors=True)
    shutil.copytree(out_dir, latest_dir)
    print(f"Saved kit: {out_dir}")


if __name__ == "__main__":
    main()
